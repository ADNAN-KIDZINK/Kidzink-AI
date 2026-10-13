"""Build an Excel workbook from a JSON description.

Usage: kz-python make_xlsx.py spec.json output.xlsx

spec.json:
{
  "currency": "AED",                       optional, used by "currency" columns
  "sheets": [
    {
      "name": "Budget",
      "columns": [
        {"header": "Item", "type": "text"},
        {"header": "Cost", "type": "currency"},
        {"header": "Due", "type": "date"}          types: text, number, integer, currency,
      ],                                                   percent, date
      "rows": [
        ["Laptops", 12000, "2026-11-01"],
        ["Total", "=SUM(B2:B2)", null]             strings starting with "=" are formulas
      ],
      "total_row": {"label": "Total", "sum": ["Cost"]}   optional: adds a SUM row
    }
  ]
}

The result is deliberately plain: bold header row, frozen header, filters, sensible
column widths and number formats. Dates are written as real dates (YYYY-MM-DD in the spec).
"""

import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

FORMATS = {
    "number": "#,##0.00",
    "integer": "#,##0",
    "percent": "0.0%",
    "date": "dd mmm yyyy",
}


def number_format(column_type, currency):
    if column_type == "currency":
        return f'#,##0.00 "{currency}"' if currency else "#,##0.00"
    return FORMATS.get(column_type)


def convert(value, column_type):
    if isinstance(value, str) and value.startswith("="):
        return value
    if column_type == "date" and isinstance(value, str) and value:
        try:
            return datetime.strptime(value[:10], "%Y-%m-%d").date()
        except ValueError:
            return value
    return value


def display_width(value):
    if value is None:
        return 0
    if isinstance(value, (date, datetime)):
        return 12
    if isinstance(value, (int, float)):
        return len(f"{value:,.2f}") + 2
    text = str(value)
    return 12 if text.startswith("=") else max(len(line) for line in text.splitlines() or [""])


def add_sheet(workbook, sheet_spec, currency, first):
    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = str(sheet_spec.get("name", "Sheet1"))[:31]
    columns = sheet_spec["columns"]
    headers = [column["header"] for column in columns]
    types = [column.get("type", "text") for column in columns]

    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    rows = sheet_spec.get("rows", [])
    for row in rows:
        sheet.append([convert(value, types[i]) if i < len(types) else value for i, value in enumerate(row)])

    last_data_row = len(rows) + 1
    total = sheet_spec.get("total_row")
    if total and rows:
        total_values = [None] * len(columns)
        total_values[0] = total.get("label", "Total")
        for name in total.get("sum", []):
            index = headers.index(name)
            letter = get_column_letter(index + 1)
            total_values[index] = f"=SUM({letter}2:{letter}{last_data_row})"
        sheet.append(total_values)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)

    for index, column_type in enumerate(types, start=1):
        fmt = number_format(column_type, currency)
        if fmt:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = fmt
        widths = [display_width(cell.value) for (cell,) in sheet.iter_rows(min_col=index, max_col=index)]
        explicit = columns[index - 1].get("width")
        sheet.column_dimensions[get_column_letter(index)].width = explicit or min(max(widths + [8]) + 2, 60)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_data_row}"


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    output = Path(sys.argv[2])
    if output.suffix.lower() != ".xlsx":
        sys.exit("The output file must end in .xlsx")
    workbook = Workbook()
    for i, sheet_spec in enumerate(spec["sheets"]):
        add_sheet(workbook, sheet_spec, spec.get("currency"), first=i == 0)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"Saved {output.resolve()}")


if __name__ == "__main__":
    main()
