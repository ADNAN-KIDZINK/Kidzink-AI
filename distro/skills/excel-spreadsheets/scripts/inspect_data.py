"""Describe a spreadsheet or CSV before working on it: sheets, columns, types, blanks, duplicates.

Usage: kz-python inspect_data.py file.xlsx|file.csv [--sheet NAME] [--rows 5]
"""

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def read_csv(path):
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return pd.read_csv(path, encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise SystemExit(f"Couldn't read {path} as text")


def describe(name, frame, sample_rows):
    print(f"## Sheet: {name}")
    print(f"{len(frame)} rows x {len(frame.columns)} columns")
    duplicates = int(frame.duplicated().sum())
    empty_rows = int(frame.isna().all(axis=1).sum())
    print(f"Fully duplicated rows: {duplicates} | Completely empty rows: {empty_rows}\n")
    print("| Column | Type | Blank | Distinct | Example |")
    print("|---|---|---|---|---|")
    for column in frame.columns:
        series = frame[column]
        non_blank = series.dropna()
        example = str(non_blank.iloc[0])[:40].replace("|", "/") if len(non_blank) else ""
        print(f"| {column} | {series.dtype} | {int(series.isna().sum())} | {series.nunique()} | {example} |")
    print(f"\nFirst {sample_rows} rows:")
    print(frame.head(sample_rows).to_string(max_colwidth=40))
    print()


def formula_note(path):
    formulas = cached = 0
    values = load_workbook(path, data_only=True, read_only=True)
    for sheet in load_workbook(path, read_only=True).worksheets:
        value_rows = values[sheet.title].iter_rows(values_only=True)
        for row, value_row in zip(sheet.iter_rows(values_only=True), value_rows):
            for cell, value in zip(row, value_row):
                if isinstance(cell, str) and cell.startswith("="):
                    formulas += 1
                    cached += value is not None
    if formulas and cached < formulas:
        print(
            f"Note: {formulas} cells hold formulas and {formulas - cached} have no saved result "
            "(the file hasn't been opened in Excel since they were written). They show as blank "
            "below; calculate them in Python if you need the numbers.\n"
        )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("file")
    parser.add_argument("--sheet")
    parser.add_argument("--rows", type=int, default=5)
    args = parser.parse_args()
    path = Path(args.file)

    if path.suffix.lower() in (".csv", ".txt"):
        describe(path.name, read_csv(path), args.rows)
        return
    sheets = pd.read_excel(path, sheet_name=args.sheet or None)
    if isinstance(sheets, pd.DataFrame):
        sheets = {args.sheet: sheets}
    print(f"Workbook {path.name}: {len(sheets)} sheet(s): {', '.join(map(str, sheets))}\n")
    formula_note(path)
    for name, frame in sheets.items():
        describe(name, frame, args.rows)


if __name__ == "__main__":
    main()
