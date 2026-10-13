"""Build a Word document from simple Markdown.

Usage: kz-python make_docx.py input.md output.docx [--title "Document title"]

Supported Markdown:
  # / ## / ### headings     paragraphs (blank line between)
  - or * bullets            two-space indent for a second level
  1. numbered items         > quotes
  | tables |                (first row is the header; the |---| line is optional)
  **bold**, *italic*        ---  (on its own line) starts a new page
Text such as [to confirm] is highlighted so reviewers can spot it.
"""

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_COLOR_INDEX
from docx.shared import Pt

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]*confirm[^\]]*\])", re.IGNORECASE)
TABLE_DIVIDER = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")


def add_inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("[") and "confirm" in part.lower():
            paragraph.add_run(part).font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            paragraph.add_run(part)


def split_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    columns = len(header)
    table = doc.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text.replace("**", ""))
        run.bold = True
    for row in body:
        cells = table.add_row().cells
        for i in range(columns):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], row[i] if i < len(row) else "")
    doc.add_paragraph()


def build(markdown, title):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    if title:
        doc.add_heading(title, level=0)

    lines = markdown.splitlines()
    # A leading "# Title" becomes the document title, so "##" sections become Heading 1.
    first_text = next((line.strip() for line in lines if line.strip()), "")
    title_from_h1 = title is None and re.match(r"^#\s", first_text) is not None
    shift = 1 if (title or title_from_h1) else 0
    paragraph_lines = []
    table_rows = []

    def flush_paragraph():
        if paragraph_lines:
            add_inline(doc.add_paragraph(), " ".join(line.strip() for line in paragraph_lines))
            paragraph_lines.clear()

    def flush_table():
        if table_rows:
            add_table(doc, table_rows)
            table_rows.clear()

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("|"):
            flush_paragraph()
            if not TABLE_DIVIDER.match(stripped):
                table_rows.append(split_row(stripped))
            continue
        flush_table()

        if not stripped:
            flush_paragraph()
            continue
        if stripped == "---":
            flush_paragraph()
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            if title_from_h1 and level == 1 and not doc.paragraphs:
                doc.add_heading(heading.group(2), level=0)
            else:
                doc.add_heading(heading.group(2), level=max(1, min(level - shift, 4)))
            continue

        bullet = re.match(r"^(\s*)[-*•]\s+(.*)$", line)
        numbered = re.match(r"^(\s*)\d+[.)]\s+(.*)$", line)
        if bullet or numbered:
            flush_paragraph()
            match = bullet or numbered
            deep = len(match.group(1).replace("\t", "  ")) >= 2
            style = ("List Bullet" if bullet else "List Number") + (" 2" if deep else "")
            add_inline(doc.add_paragraph(style=style), match.group(2))
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            add_inline(doc.add_paragraph(style="Quote"), stripped.lstrip("> ").strip())
            continue

        paragraph_lines.append(stripped)

    flush_paragraph()
    flush_table()
    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input", help="Markdown file")
    parser.add_argument("output", help="Word file to create (.docx)")
    parser.add_argument("--title", help="Title shown at the top of the document")
    args = parser.parse_args()

    output = Path(args.output)
    if output.suffix.lower() != ".docx":
        sys.exit("The output file must end in .docx")
    doc = build(Path(args.input).read_text(encoding="utf-8"), args.title)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Saved {output.resolve()}")


if __name__ == "__main__":
    main()
