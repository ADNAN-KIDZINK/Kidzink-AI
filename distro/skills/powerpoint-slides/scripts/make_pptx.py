"""Build a PowerPoint deck (16:9) from a JSON outline.

Usage: kz-python make_pptx.py outline.json output.pptx

outline.json:
{
  "title": "Cybersecurity plan",            first slide
  "subtitle": "Leadership review, Oct 2026", optional
  "slides": [
    {"type": "section", "title": "Where we are"},
    {"type": "bullets", "title": "Gaps", "bullets": ["Email", ["sub-point"], "Backups"],
     "notes": "Speaker notes"},
    {"type": "two_column", "title": "Options", "left_heading": "Option A", "left": ["..."],
     "right_heading": "Option B", "right": ["..."]},
    {"type": "table", "title": "Costs", "columns": ["Item", "AED"], "rows": [["VAPT", "16,000"]]},
    {"type": "closing", "title": "Next steps", "bullets": ["..."]}
  ]
}
A nested list inside "bullets" makes second-level points for the item before it.
Keep to about 6 bullets per slide; split longer content across slides.
"""

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

TITLE, TITLE_AND_CONTENT, SECTION, TWO_CONTENT, TITLE_ONLY = 0, 1, 2, 3, 5


def fill_bullets(frame, items):
    frame.clear()
    first = True
    for item in items:
        children = item if isinstance(item, list) else None
        texts = children if children else [item]
        for text in texts:
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            paragraph.text = str(text)
            paragraph.level = 1 if children else 0
            first = False


def add_notes(slide, spec):
    if spec.get("notes"):
        slide.notes_slide.notes_text_frame.text = spec["notes"]


def bullets_slide(deck, spec):
    slide = deck.slides.add_slide(deck.slide_layouts[TITLE_AND_CONTENT])
    slide.shapes.title.text = spec.get("title", "")
    fill_bullets(slide.placeholders[1].text_frame, spec.get("bullets", []))
    return slide


def two_column_slide(deck, spec):
    slide = deck.slides.add_slide(deck.slide_layouts[TWO_CONTENT])
    slide.shapes.title.text = spec.get("title", "")
    for placeholder, side in ((slide.placeholders[1], "left"), (slide.placeholders[2], "right")):
        heading = spec.get(f"{side}_heading")
        items = ([heading] if heading else []) + list(spec.get(side, []))
        fill_bullets(placeholder.text_frame, items)
        if heading:
            placeholder.text_frame.paragraphs[0].runs[0].font.bold = True
    return slide


def table_slide(deck, spec):
    slide = deck.slides.add_slide(deck.slide_layouts[TITLE_ONLY])
    slide.shapes.title.text = spec.get("title", "")
    columns, rows = spec["columns"], spec.get("rows", [])
    left, top = Inches(0.6), Inches(1.6)
    width = deck.slide_width - 2 * left
    height = Emu(int(Inches(0.45)) * (len(rows) + 1))
    table = slide.shapes.add_table(len(rows) + 1, len(columns), left, top, width, height).table
    for c, header in enumerate(columns):
        table.cell(0, c).text = str(header)
    for r, row in enumerate(rows, start=1):
        for c in range(len(columns)):
            table.cell(r, c).text = str(row[c]) if c < len(row) and row[c] is not None else ""
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(14)
    return slide


def section_slide(deck, spec):
    slide = deck.slides.add_slide(deck.slide_layouts[SECTION])
    slide.shapes.title.text = spec.get("title", "")
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = spec.get("subtitle", "")
    return slide


def widen_to_16_9(deck):
    # The default template is 4:3; its master and layout placeholders must be stretched too,
    # or every slide's content sits in the left three quarters. Layouts go first: a layout
    # placeholder without its own position reads the master's, which must not be scaled yet.
    ratio = Inches(13.333) / deck.slide_width
    deck.slide_width = Inches(13.333)
    for container in [*deck.slide_layouts, deck.slide_master]:
        for shape in container.shapes:
            shape.left = int(shape.left * ratio)
            shape.width = int(shape.width * ratio)


BUILDERS = {
    "bullets": bullets_slide,
    "closing": bullets_slide,
    "two_column": two_column_slide,
    "table": table_slide,
    "section": section_slide,
}


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    outline = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    output = Path(sys.argv[2])
    if output.suffix.lower() != ".pptx":
        sys.exit("The output file must end in .pptx")

    deck = Presentation()
    widen_to_16_9(deck)

    title_slide = deck.slides.add_slide(deck.slide_layouts[TITLE])
    title_slide.shapes.title.text = outline.get("title", "")
    title_slide.placeholders[1].text = outline.get("subtitle", "")

    for spec in outline.get("slides", []):
        builder = BUILDERS.get(spec.get("type", "bullets"))
        if builder is None:
            sys.exit(f"Unknown slide type {spec.get('type')!r}; use one of {', '.join(BUILDERS)}")
        add_notes(builder(deck, spec), spec)

    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
    print(f"Saved {output.resolve()} ({len(deck.slides)} slides)")


if __name__ == "__main__":
    main()
